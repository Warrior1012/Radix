"""Upload handling for JD/resume documents with layout-aware text extraction."""
from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from fastapi import HTTPException
from pypdf import PdfReader

try:  # PyMuPDF's geometric reading order is better for many multi-column PDFs.
    import fitz  # type: ignore
except ImportError:  # Keep source usable if a deployment has not installed the optional parser.
    fitz = None

MAX_TEXT_CHARS = 40000
MIN_EXTRACTED_CHARS = 40


def safe_filename(name: str) -> str:
    return (name or "upload").replace("\\", "/").split("/")[-1]


def _pdf_text(content: bytes) -> tuple[str, str]:
    """Return (text, extraction_method). Prefer PyMuPDF sorted reading order, then pypdf."""
    errors: list[str] = []
    if fitz is not None:
        try:
            with fitz.open(stream=content, filetype="pdf") as pdf:
                pages = [page.get_text("text", sort=True) or "" for page in pdf]
                text = "\n".join(pages).strip()
            if text:
                return text, "pymupdf_sorted_text"
        except Exception as exc:
            errors.append(type(exc).__name__)
    try:
        text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages).strip()
        if text:
            return text, "pypdf_text"
    except Exception as exc:
        errors.append(type(exc).__name__)
    if errors:
        raise HTTPException(400, f"Could not read this PDF ({', '.join(dict.fromkeys(errors))}).")
    return "", "no_text"


def extract_text_with_metadata(filename: str, content: bytes) -> tuple[str, dict[str, object]]:
    ext = Path(filename).suffix.lower()
    method = ""
    try:
        if ext == ".pdf":
            text, method = _pdf_text(content)
        elif ext == ".docx":
            d = Document(io.BytesIO(content))
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            for table in d.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            text = "\n".join(parts).strip()
            method = "python_docx_paragraphs_tables"
        else:
            raise HTTPException(415, "Only PDF and DOCX files are supported.")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(400, f"Could not read this document: {type(exc).__name__}") from exc
    text = text.strip()
    if len(text) < MIN_EXTRACTED_CHARS:
        if ext == ".pdf":
            raise HTTPException(422, "Very little text was extracted from this PDF. It may be scanned/image-only or encrypted. This build supports text PDFs and DOCX; OCR/direct image-PDF extraction is not enabled. Use a text-based PDF or DOCX for this run.")
        raise HTTPException(422, "Very little text was extracted from this document. Check that the file contains selectable text.")
    return text[:MAX_TEXT_CHARS], {"method": method, "characters": min(len(text), MAX_TEXT_CHARS), "truncated": len(text) > MAX_TEXT_CHARS}


def extract_text(filename: str, content: bytes) -> str:
    """Backward-compatible text-only wrapper used by existing callers/tests."""
    text, _ = extract_text_with_metadata(filename, content)
    return text
